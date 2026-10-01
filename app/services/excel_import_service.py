import math
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, StrictBool, ValidationError
from sqlalchemy.orm import Session


from app.models import InboundEvent, ProcessingJob
from app.schemas import ExcelImportAccepted, ExcelImportStatus, ExcelRowError
from app.services.product_processor import ProductProcessor
from app.services.sync_service import SyncService

HEADER = ("product_id", "sku", "name", "is_active", "source_version")
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 10_000
ERROR_STATES = ("VALIDATION_ERROR", "FAILED", "CONFLICT")

class ExcelProductRow(BaseModel):
    product_id: int = Field(gt=0, le=2_147_483_647, strict = True)
    sku: str | None = Field(max_length=100, strict=True)
    name: str | None = Field(max_length=255, strict=True)
    is_active: StrictBool
    source_version: int = Field(gt=0, le=9_223_372_036_854_775_807, strict=True)
    
    model_config = ConfigDict(extra="forbid")
    
@dataclass(frozen=True)
class ParsedRow:
    item_index: int
    raw_payload: dict
    payload: dict | None
    product_id: int | None
    source_version: int | None
    error_message: str | None
    
def _json_safe(value: object) -> tuple[object,bool]:
    if value is None or isinstance(value, (str,bool,int)):
        return value, False
    if isinstance(value, (datetime,date)):
        return value.isoformat(), False
    if isinstance(value, Decimal):
        return str(value), not value.is_finite()
    if isinstance(value, float):
        return (value, False) if math.isfinite(value) else (None, True)
    try:
        return str(value), True
    except Exception:
        return f"<unserializable {type(value).__name__}>", True
    
def _parse_row(index: int, cells: tuple) -> ParsedRow:
    raw = {}
    conversion_errors = []
    padded_cells = tuple(cells[:len(HEADER)]) + (None,) * max(0,len(HEADER) - len(cells))
    for key, value in zip(HEADER, padded_cells):
        raw[key], invalid = _json_safe(value)
        if invalid:
            conversion_errors.append(f"{key}: unsupported or non-finite cell value")
            
    product_id = raw["product_id"]
    if type(product_id) is not int or not 0 < product_id <= 2_147_483_647:
        product_id = None
    source_version = raw["source_version"]
    if type(source_version) is not int or not 0 < source_version <= 9_223_372_036_854_775_807:
        source_version = None
        
    values = dict(zip(HEADER, padded_cells))
    active = values["is_active"]
    if active == "true":
        values["is_active"] = True
    elif active == "false":
        values["is_active"] = False
    try:
        parsed = ExcelProductRow.model_validate(values)
        if conversion_errors:
            raise ValueError("; ".join(conversion_errors))
    except ValidationError as exc:
        errors = [f"{'.'.join(map(str, err['loc']))}: {err['msg']}" for err in exc.errors()]
        return ParsedRow(index, raw, None, product_id, source_version,
                         "; ".join(conversion_errors + errors)[:1000])
    except ValueError as exc:
        return ParsedRow(index, raw, None, product_id, source_version, str(exc)[:1000])

    return ParsedRow(index, raw, parsed.model_dump(exclude={"source_version"}),
                     parsed.product_id, parsed.source_version, None)

def parse_workbook(contents: bytes) -> list[ParsedRow]:
    if len(contents) > MAX_FILE_BYTES:
        raise HTTPException(413, "Excel file exceeds 5 MiB")
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        try:
            if "Products" not in workbook.sheetnames:
                raise HTTPException(422, "Workbook must contain a Products sheet")
            sheet = workbook["Products"]
            iterator = sheet.iter_rows(values_only=True)
            header = tuple(next(iterator, ()))
            if header != HEADER:
                raise HTTPException(422, f"Products header must be: {', '.join(HEADER)}")
            rows = []
            for physical_index, cells in enumerate(iterator, start=2):
                if physical_index - 1 > MAX_DATA_ROWS:
                    raise HTTPException(413, "Excel sheet exceeds 10000 data rows")
                # Rows produced by openpyxl are padded only as far as its sheet dimension.
                rows.append(_parse_row(physical_index, cells[:len(HEADER)]))
            return rows
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Excel workbook is damaged or unreadable") from exc


class ExcelImportService:
    @staticmethod
    def create(db: Session, filename: str, rows: list[ParsedRow]) -> ProcessingJob:
        job = ProcessingJob(
            job_id=uuid.uuid4(), job_type="EXCEL", status="PENDING",
            parameters={}, file_name=filename, input_complete=True,
        )
        db.add(job)
        db.flush()
        now = datetime.now(timezone.utc)
        for row in rows:
            invalid = row.error_message is not None
            db.add(InboundEvent(
                event_id=uuid.uuid4(), channel="EXCEL", job_id=job.job_id,
                item_index=row.item_index, external_event_id=None,
                product_id=row.product_id, source_version=row.source_version,
                raw_payload=row.raw_payload, payload=row.payload,
                status="VALIDATION_ERROR" if invalid else "RECEIVED",
                error_code="VALIDATION_ERROR" if invalid else None,
                error_message=row.error_message,
                processed_at=now if invalid else None,
                attempt_count=0,
            ))
        try:
            db.commit()
        except Exception:
            db.rollback()
            raise
        db.refresh(job)
        return job

    @staticmethod
    def get(db: Session, batch_id: uuid.UUID) -> ProcessingJob:
        job = db.get(ProcessingJob, batch_id)
        if job is None or job.job_type != "EXCEL":
            raise HTTPException(404, "Excel import not found")
        return job

    @staticmethod
    def accepted(job: ProcessingJob, count: int) -> ExcelImportAccepted:
        return ExcelImportAccepted(
            batch_id=str(job.job_id), job_type="EXCEL", status="PENDING",
            total_count=count, status_url=f"/v1/imports/{job.job_id}",
        )

    @staticmethod
    def response(db: Session, job: ProcessingJob) -> ExcelImportStatus:
        return ExcelImportStatus(
            batch_id=str(job.job_id), job_type="EXCEL", status=job.status,
            file_name=job.file_name, input_complete=job.input_complete,
            **SyncService.counters(db, job.job_id),
        )

    @staticmethod
    def errors(db: Session, job: ProcessingJob) -> list[ExcelRowError]:
        records = (db.query(InboundEvent)
                   .filter(InboundEvent.job_id == job.job_id,
                           InboundEvent.status.in_(ERROR_STATES))
                   .order_by(InboundEvent.item_index).all())
        return [ExcelRowError(
            item_index=event.item_index, status=event.status,
            error_code=event.error_code, error_message=event.error_message,
            raw_payload=event.raw_payload,
        ) for event in records]


def process_excel_job(db: Session, job: ProcessingJob) -> None:
    """Apply remaining RECEIVED rows; completed rows survive a restart."""
    if job.job_type != "EXCEL" or not job.input_complete:
        return
    if job.status not in ("PENDING", "PROCESSING", "RETRYING"):
        return
    job.status = "PROCESSING"
    job.started_at = datetime.now(timezone.utc)
    job.attempt_count += 1
    db.commit()
    event_ids = [event_id for (event_id,) in (
        db.query(InboundEvent.event_id)
        .filter(InboundEvent.job_id == job.job_id,
                InboundEvent.status.in_(("RECEIVED", "RETRYING")))
        .order_by(InboundEvent.item_index).all()
    )]
    for event_id in event_ids:
        ProductProcessor.process(db, event_id)
        job.started_at = datetime.now(timezone.utc)
        db.commit()
    SyncService.finalize(db, job)


def process_excel_in_background(batch_id: uuid.UUID) -> None:
    from app.core.db import SessionLocal

    with SessionLocal() as db:
        job = db.get(ProcessingJob, batch_id, with_for_update=True)
        if job is not None and job.status == "PENDING":
            process_excel_job(db, job)
